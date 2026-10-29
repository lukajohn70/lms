#!/usr/bin/env python3
"""
DLHS Portal Broadsheet Downloader & Compiler (2025/2026 Session)
Downloads broadsheet data from dlhsportal.org for:
  - SS 1 Science (Merged Science A & Science B)
  - SS 2 Humanities
  - SS 2 Science (Science A)
  - SS 1 Humanities (for safety & completeness)
Generates:
  1. Professional Excel Workbooks (.xlsx) with styled banners, zebra rows, positions, averages, and statistics
  2. Standard CSV files
  3. Portal-style HTML-based .xls files
Saves outputs to:
  - /Users/john/Downloads/DLHS_Broadsheets_2025_2026/
  - /Applications/MAMP/htdocs/lms/downloads/broadsheets_2025_2026/
"""

import os
import sys
import json
import re
import shutil
import ssl
import urllib.request
import urllib.parse
import http.cookiejar
from collections import defaultdict
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# --- Configuration ---
PORTAL_LOGIN_URL = 'https://dlhsportal.org/sms/login.php'
PORTAL_EMAIL     = 'ictkaduna@deeperlifehighschool.org'
PORTAL_PASSWORD  = '12'
SECTION_ID       = '13'  # 2025/2026 Academic Session

DEST_DIRS = [
    '/Users/john/Downloads/DLHS_Broadsheets_2025_2026',
    '/Applications/MAMP/htdocs/lms/downloads/broadsheets_2025_2026'
]

TERMS = [
    ('term1', 'First Term'),
    ('term2', 'Second Term'),
    ('term3', 'Third Term')
]

CLASSES = {
    'ss1_sci_a': {'cid': '502', 'name': 'SS 1 SCIENCE A', 'label': 'SS 1 Science A'},
    'ss1_sci_b': {'cid': '862', 'name': 'SS 1 SCIENCE B', 'label': 'SS 1 Science B'},
    'ss2_hum':   {'cid': '473', 'name': 'SS 2 HUMANITIES', 'label': 'SS 2 Humanities'},
    'ss2_sci_a': {'cid': '501', 'name': 'SS 2 SCIENCE A', 'label': 'SS 2 Science A'},
    'ss2_sci_b': {'cid': '444', 'name': 'SS 2 SCIENCE B', 'label': 'SS 2 Science B'},
    'ss1_hum':   {'cid': '472', 'name': 'SS 1 HUMANITIES', 'label': 'SS 1 Humanities'},
}

# --- Establish Authenticated Portal Session ---
ctx = ssl.create_default_context()
ctx.check_hostname = False
ctx.verify_mode = ssl.CERT_NONE

cj = http.cookiejar.CookieJar()
opener = urllib.request.build_opener(
    urllib.request.HTTPCookieProcessor(cj),
    urllib.request.HTTPSHandler(context=ctx)
)

print("🔑 Logging into dlhsportal.org as ICT Admin...")
login_payload = urllib.parse.urlencode({
    'email': PORTAL_EMAIL,
    'password': PORTAL_PASSWORD,
    'user_type': 'admin',
    'submit': 'yes'
}).encode()

req = urllib.request.Request(
    PORTAL_LOGIN_URL,
    data=login_payload,
    headers={'User-Agent': 'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7)'}
)
res = opener.open(req)
login_html = res.read().decode('utf-8', errors='ignore')

if 'logout' in login_html.lower() or 'dashboard' in login_html.lower():
    print("✅ Successfully authenticated to DLHS Portal.")
else:
    print("❌ Login failed. Response snippet:")
    print(login_html[:500])
    sys.exit(1)

# --- Fetch Data for all Classes and Terms ---
raw_cache = {}

def fetch_broadsheet(class_id, term_code):
    cache_key = f"{class_id}_{term_code}"
    if cache_key in raw_cache:
        return raw_cache[cache_key]
    
    url = f"https://dlhsportal.org/sms/admin/fetch_broadsheet_previous.php?class_ID={class_id}&section_ID={SECTION_ID}&term={term_code}"
    req = urllib.request.Request(
        url,
        headers={'User-Agent': 'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7)'}
    )
    resp = opener.open(req)
    data = json.loads(resp.read().decode('utf-8', errors='ignore'))
    raw_cache[cache_key] = data
    return data

print("\n📥 Fetching all broadsheet records for 2025/2026 Academic Session...")
for ckey, cinfo in CLASSES.items():
    for term_code, term_name in TERMS:
        data = fetch_broadsheet(cinfo['cid'], term_code)
        num_students = len(data.get('students', []))
        num_subjects = len(data.get('subjects', []))
        print(f"  • {cinfo['label']:16} | {term_name:12}: {num_students:2} students, {num_subjects:2} subjects")

def clean_subj_name(name):
    clean = re.sub(r'\s+', ' ', name).strip().upper()
    return clean

def ordinal(n):
    if 11 <= (n % 100) <= 13:
        suffix = 'th'
    else:
        suffix = {1: 'st', 2: 'nd', 3: 'rd'}.get(n % 10, 'th')
    return f"{n}{suffix}"

# --- Helper to format Student Tables ---
def prepare_class_table(raw_data, arm_label=None):
    """Parses raw JSON from fetch_broadsheet_previous into a structured table."""
    subjects_raw = raw_data.get('subjects', [])
    students_raw = raw_data.get('students', [])
    
    # Map subject ID -> Clean Name
    sub_map = {}
    ordered_subs = []
    seen_subs = set()
    for s in subjects_raw:
        sname = clean_subj_name(s['name'])
        sub_map[str(s['id'])] = sname
        if sname not in seen_subs:
            seen_subs.add(sname)
            ordered_subs.append(sname)
    
    records = []
    for s in students_raw:
        sname = s.get('name', '').strip().upper()
        scores_by_id = s.get('scores', {})
        # Map scores to subject names
        sub_scores = {}
        for sid, score in scores_by_id.items():
            s_name = sub_map.get(str(sid))
            if s_name:
                val = float(score) if score is not None else 0.0
                sub_scores[s_name] = val
        
        # Calculate statistics
        total = float(s.get('total', 0.0))
        offered = [sc for sc in sub_scores.values() if sc > 0]
        avg = round(total / len(offered), 2) if offered else 0.0
        
        records.append({
            'name': sname,
            'arm': arm_label or '',
            'scores': sub_scores,
            'total': total,
            'avg': avg,
            'subjects_offered': len(offered)
        })
    
    # Sort by total descending
    records.sort(key=lambda x: x['total'], reverse=True)
    for idx, r in enumerate(records):
        r['rank'] = idx + 1
        r['position'] = ordinal(idx + 1)
        
    return ordered_subs, records

def prepare_merged_ss1_science(term_code):
    """Merges SS 1 Science A and SS 1 Science B into a single cohesive broadsheet."""
    data_a = fetch_broadsheet(CLASSES['ss1_sci_a']['cid'], term_code)
    data_b = fetch_broadsheet(CLASSES['ss1_sci_b']['cid'], term_code)
    
    _, records_a = prepare_class_table(data_a, arm_label='Science A')
    _, records_b = prepare_class_table(data_b, arm_label='Science B')
    
    # Union of all subjects, with canonical core subjects first
    core_order = [
        'MATHEMATICS', 'ENGLISH LANGUAGE', 'PHYSICS', 'CHEMISTRY', 'BIOLOGY',
        'FURTHER MATHS', 'AGRICULTURAL SCIENCE', 'CIVIC EDUCATION', 'DATA PROCESSING',
        'COMPUTER STUDIES', 'ROBOTICS AND A.I.', 'TECHNICAL DRAWING', 'GEOGRAPHY',
        'ECONOMICS', 'FRENCH', 'LITERATURE IN ENGLISH', 'ANIMAL HUSBANDRY',
        'FOOD AND NUTRITION', 'CATERING CRAFT PRACTICE', 'VISUAL ART',
        'CHRISTIAN RELIGIOUS STUDIES', 'GOVERNMENT', 'LET',
        'FASHION DESIGN AND GARMENT MAKING', 'SOLAR PHOTOVOLTAIC INSTALLATION AND MAINTENANCE'
    ]
    
    all_seen = set()
    for s in data_a.get('subjects', []) + data_b.get('subjects', []):
        all_seen.add(clean_subj_name(s['name']))
        
    ordered_subs = [s for s in core_order if s in all_seen]
    for s in sorted(all_seen):
        if s not in ordered_subs:
            ordered_subs.append(s)
            
    all_records = records_a + records_b
    all_records.sort(key=lambda x: x['total'], reverse=True)
    
    for idx, r in enumerate(all_records):
        r['rank'] = idx + 1
        r['position'] = ordinal(idx + 1)
        
    return ordered_subs, all_records

# --- Styling helpers for openpyxl ---
FONT_FAMILY = "Calibri"
NAVY_HEADER = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
BLUE_SUBHEADER = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
TABLE_HEADER_FILL = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
ZEBRA_FILL = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
WHITE_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
TOTAL_FILL = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
AVG_FILL   = PatternFill(start_color="E0E7FF", end_color="E0E7FF", fill_type="solid")
RANK_FILL  = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
STATS_LABEL_FILL = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

TITLE_FONT = Font(name=FONT_FAMILY, size=15, bold=True, color="FFFFFF")
SUBTITLE_FONT = Font(name=FONT_FAMILY, size=11, italic=True, color="FFFFFF")
CLASS_TITLE_FONT = Font(name=FONT_FAMILY, size=12, bold=True, color="FFFFFF")
TH_FONT = Font(name=FONT_FAMILY, size=9, bold=True, color="FFFFFF")
TD_FONT = Font(name=FONT_FAMILY, size=10, color="0F172A")
TD_BOLD = Font(name=FONT_FAMILY, size=10, bold=True, color="0F172A")
STATS_FONT = Font(name=FONT_FAMILY, size=10, bold=True, color="1E293B")

THIN_BORDER_COLOR = "CBD5E1"
thin_side = Side(style='thin', color=THIN_BORDER_COLOR)
CELL_BORDER = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
DOUBLE_BOTTOM = Border(left=thin_side, right=thin_side, top=thin_side, bottom=Side(style='double', color="0F172A"))

def write_broadsheet_sheet(ws, title_class, term_label, subjects, records, include_arm=False):
    """Writes a polished, professional broadsheet to an openpyxl Worksheet."""
    ws.views.sheetView[0].showGridLines = True
    
    # Calculate total columns
    # Fixed cols: S/N (A), [Arm (B) if include_arm], Name (B/C)
    # Followed by subjects
    # Followed by: Subj Offered, Total Score, Average, Position
    start_col = 1
    total_cols = (3 if include_arm else 2) + len(subjects) + 4
    
    # Row 1: School Banner
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
    cell_r1 = ws.cell(row=1, column=1, value="DEEPER LIFE HIGH SCHOOL, KADUNA CAMPUS")
    cell_r1.font = TITLE_FONT
    cell_r1.fill = NAVY_HEADER
    cell_r1.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28
    
    # Row 2: Sub-banner
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=total_cols)
    cell_r2 = ws.cell(row=2, column=1, value=f"OFFICIAL BROADSHEET — 2025/2026 ACADEMIC SESSION — {term_label.upper()}")
    cell_r2.font = SUBTITLE_FONT
    cell_r2.fill = NAVY_HEADER
    cell_r2.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20
    
    # Row 3: Class Header
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=total_cols)
    cell_r3 = ws.cell(row=3, column=1, value=f"CLASS: {title_class.upper()} | ENROLLMENT: {len(records)} STUDENTS")
    cell_r3.font = CLASS_TITLE_FONT
    cell_r3.fill = BLUE_SUBHEADER
    cell_r3.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[3].height = 22
    
    # Row 4: Column Headers
    headers = ["S/N"]
    if include_arm:
        headers.append("Arm")
    headers.append("Student Name")
    headers.extend(subjects)
    headers.extend(["Subjects Offered", "Total Score", "Average (%)", "Position"])
    
    ws.row_dimensions[4].height = 36
    for c_idx, htext in enumerate(headers, start=1):
        c = ws.cell(row=4, column=c_idx, value=htext)
        c.font = TH_FONT
        c.fill = TABLE_HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = CELL_BORDER

    # Rows 5+: Student Records
    data_start_row = 5
    for r_idx, rec in enumerate(records, start=data_start_row):
        is_even = (r_idx % 2 == 0)
        row_fill = ZEBRA_FILL if is_even else WHITE_FILL
        ws.row_dimensions[r_idx].height = 20
        
        curr_c = 1
        # S/N
        c = ws.cell(row=r_idx, column=curr_c, value=rec['rank'])
        c.font = TD_FONT
        c.fill = row_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = CELL_BORDER
        curr_c += 1
        
        # Arm (if merged)
        if include_arm:
            c = ws.cell(row=r_idx, column=curr_c, value=rec['arm'])
            c.font = TD_FONT
            c.fill = row_fill
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = CELL_BORDER
            curr_c += 1
            
        # Student Name
        c = ws.cell(row=r_idx, column=curr_c, value=rec['name'])
        c.font = TD_BOLD
        c.fill = row_fill
        c.alignment = Alignment(horizontal="left", vertical="center")
        c.border = CELL_BORDER
        curr_c += 1
        
        # Subject Scores
        for sub in subjects:
            sc = rec['scores'].get(sub, 0.0)
            c = ws.cell(row=r_idx, column=curr_c)
            if sc > 0:
                c.value = sc
                c.number_format = '0.0' if (sc % 1 != 0) else '0'
                c.font = TD_FONT
            else:
                c.value = "-"
                c.font = Font(name=FONT_FAMILY, size=9, color="94A3B8")
            c.fill = row_fill
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = CELL_BORDER
            curr_c += 1
            
        # Subjects Offered
        c = ws.cell(row=r_idx, column=curr_c, value=rec['subjects_offered'])
        c.font = TD_FONT
        c.fill = row_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = CELL_BORDER
        curr_c += 1
        
        # Total Score
        c = ws.cell(row=r_idx, column=curr_c, value=rec['total'])
        c.font = TD_BOLD
        c.fill = TOTAL_FILL
        c.number_format = '0.0' if (rec['total'] % 1 != 0) else '0'
        c.alignment = Alignment(horizontal="right", vertical="center")
        c.border = CELL_BORDER
        curr_c += 1
        
        # Average
        c = ws.cell(row=r_idx, column=curr_c, value=rec['avg'])
        c.font = TD_BOLD
        c.fill = AVG_FILL
        c.number_format = '0.00'
        c.alignment = Alignment(horizontal="right", vertical="center")
        c.border = CELL_BORDER
        curr_c += 1
        
        # Position
        c = ws.cell(row=r_idx, column=curr_c, value=rec['position'])
        c.font = TD_BOLD
        c.fill = RANK_FILL
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = CELL_BORDER

    # Summary Statistics rows
    last_data_row = data_start_row + len(records) - 1
    if len(records) > 0:
        stat_rows = [
            ("CLASS AVERAGE", f"=AVERAGE({get_column_letter(total_cols-1)}{data_start_row}:{get_column_letter(total_cols-1)}{last_data_row})", "0.00"),
            ("HIGHEST TOTAL SCORE", f"=MAX({get_column_letter(total_cols-2)}{data_start_row}:{get_column_letter(total_cols-2)}{last_data_row})", "0.0"),
            ("LOWEST TOTAL SCORE", f"=MIN({get_column_letter(total_cols-2)}{data_start_row}:{get_column_letter(total_cols-2)}{last_data_row})", "0.0")
        ]
        
        name_col_idx = 3 if include_arm else 2
        for s_idx, (label, formula, num_fmt) in enumerate(stat_rows, start=last_data_row + 2):
            ws.row_dimensions[s_idx].height = 20
            # Label
            ws.merge_cells(start_row=s_idx, start_column=1, end_row=s_idx, end_column=name_col_idx)
            lbl_cell = ws.cell(row=s_idx, column=1, value=label)
            lbl_cell.font = STATS_FONT
            lbl_cell.fill = STATS_LABEL_FILL
            lbl_cell.alignment = Alignment(horizontal="right", vertical="center")
            lbl_cell.border = CELL_BORDER
            
            # Put formula in Total / Avg column
            val_col = (total_cols - 1) if "AVERAGE" in label else (total_cols - 2)
            c = ws.cell(row=s_idx, column=val_col, value=formula)
            c.font = STATS_FONT
            c.fill = STATS_LABEL_FILL
            c.number_format = num_fmt
            c.alignment = Alignment(horizontal="right", vertical="center")
            c.border = DOUBLE_BOTTOM

    # Freeze panes at student data start
    ws.freeze_panes = ws.cell(row=5, column=1)
    
    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            # Skip title merged rows for width calculation
            if cell.row in [1, 2, 3]:
                continue
            if cell.value:
                val_str = str(cell.value)
                max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 11)
        
    # Give student name and S/N specific generous widths
    ws.column_dimensions['A'].width = 6
    if include_arm:
        ws.column_dimensions['B'].width = 12
        ws.column_dimensions['C'].width = 32
    else:
        ws.column_dimensions['B'].width = 32

def write_annual_summary_sheet(ws, title_class, term_records_dict, include_arm=False):
    """Creates an Annual Cumulative summary sheet aggregating Term 1, Term 2, and Term 3."""
    ws.views.sheetView[0].showGridLines = True
    
    # Gather all unique students across terms
    all_students = {}
    for t_code, (_, recs) in term_records_dict.items():
        for r in recs:
            sname = r['name']
            if sname not in all_students:
                all_students[sname] = {
                    'name': sname,
                    'arm': r['arm'],
                    'terms': {}
                }
            all_students[sname]['terms'][t_code] = r
            
    summary_list = []
    for sname, sdata in all_students.items():
        t1 = sdata['terms'].get('term1', {})
        t2 = sdata['terms'].get('term2', {})
        t3 = sdata['terms'].get('term3', {})
        
        t1_tot = t1.get('total', 0.0)
        t2_tot = t2.get('total', 0.0)
        t3_tot = t3.get('total', 0.0)
        
        t1_avg = t1.get('avg', 0.0)
        t2_avg = t2.get('avg', 0.0)
        t3_avg = t3.get('avg', 0.0)
        
        valid_avgs = [a for a in [t1_avg, t2_avg, t3_avg] if a > 0]
        cum_avg = round(sum(valid_avgs) / len(valid_avgs), 2) if valid_avgs else 0.0
        cum_tot = round(t1_tot + t2_tot + t3_tot, 2)
        
        summary_list.append({
            'name': sname,
            'arm': sdata['arm'],
            't1_tot': t1_tot,
            't1_avg': t1_avg,
            't2_tot': t2_tot,
            't2_avg': t2_avg,
            't3_tot': t3_tot,
            't3_avg': t3_avg,
            'cum_tot': cum_tot,
            'cum_avg': cum_avg,
        })
        
    summary_list.sort(key=lambda x: x['cum_avg'], reverse=True)
    for idx, r in enumerate(summary_list):
        r['rank'] = idx + 1
        r['position'] = ordinal(idx + 1)
        
    total_cols = (3 if include_arm else 2) + 9
    
    # Header Banners
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
    c1 = ws.cell(row=1, column=1, value="DEEPER LIFE HIGH SCHOOL, KADUNA CAMPUS")
    c1.font = TITLE_FONT
    c1.fill = NAVY_HEADER
    c1.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28
    
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=total_cols)
    c2 = ws.cell(row=2, column=1, value="ANNUAL CUMULATIVE BROADSHEET — 2025/2026 ACADEMIC SESSION (ALL TERMS)")
    c2.font = SUBTITLE_FONT
    c2.fill = NAVY_HEADER
    c2.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20
    
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=total_cols)
    c3 = ws.cell(row=3, column=1, value=f"CLASS: {title_class.upper()} | TOTAL ENROLLMENT: {len(summary_list)} STUDENTS")
    c3.font = CLASS_TITLE_FONT
    c3.fill = BLUE_SUBHEADER
    c3.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[3].height = 22
    
    headers = ["S/N"]
    if include_arm:
        headers.append("Arm")
    headers.extend([
        "Student Name",
        "Term 1 Total", "Term 1 Avg (%)",
        "Term 2 Total", "Term 2 Avg (%)",
        "Term 3 Total", "Term 3 Avg (%)",
        "Cumulative Total", "Cumulative Avg (%)",
        "Annual Position"
    ])
    
    ws.row_dimensions[4].height = 32
    for c_idx, htext in enumerate(headers, start=1):
        c = ws.cell(row=4, column=c_idx, value=htext)
        c.font = TH_FONT
        c.fill = TABLE_HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = CELL_BORDER

    data_start_row = 5
    for r_idx, rec in enumerate(summary_list, start=data_start_row):
        is_even = (r_idx % 2 == 0)
        row_fill = ZEBRA_FILL if is_even else WHITE_FILL
        ws.row_dimensions[r_idx].height = 20
        
        curr_c = 1
        # S/N
        c = ws.cell(row=r_idx, column=curr_c, value=rec['rank'])
        c.font = TD_FONT
        c.fill = row_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = CELL_BORDER
        curr_c += 1
        
        # Arm
        if include_arm:
            c = ws.cell(row=r_idx, column=curr_c, value=rec['arm'])
            c.font = TD_FONT
            c.fill = row_fill
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = CELL_BORDER
            curr_c += 1
            
        # Name
        c = ws.cell(row=r_idx, column=curr_c, value=rec['name'])
        c.font = TD_BOLD
        c.fill = row_fill
        c.alignment = Alignment(horizontal="left", vertical="center")
        c.border = CELL_BORDER
        curr_c += 1
        
        # Term values
        for val, fmt in [
            (rec['t1_tot'], '0.0'), (rec['t1_avg'], '0.00'),
            (rec['t2_tot'], '0.0'), (rec['t2_avg'], '0.00'),
            (rec['t3_tot'], '0.0'), (rec['t3_avg'], '0.00'),
            (rec['cum_tot'], '0.0'), (rec['cum_avg'], '0.00')
        ]:
            c = ws.cell(row=r_idx, column=curr_c, value=val if val > 0 else "-")
            c.font = TD_BOLD if "cum" in fmt or "Total" in str(val) else TD_FONT
            c.fill = TOTAL_FILL if fmt == '0.0' and "cum" in str(rec) else (AVG_FILL if fmt == '0.00' else row_fill)
            if val > 0:
                c.number_format = fmt
            c.alignment = Alignment(horizontal="right", vertical="center")
            c.border = CELL_BORDER
            curr_c += 1
            
        # Position
        c = ws.cell(row=r_idx, column=curr_c, value=rec['position'])
        c.font = TD_BOLD
        c.fill = RANK_FILL
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = CELL_BORDER

    ws.freeze_panes = ws.cell(row=5, column=1)
    
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row in [1, 2, 3]:
                continue
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
        
    ws.column_dimensions['A'].width = 6
    if include_arm:
        ws.column_dimensions['B'].width = 12
        ws.column_dimensions['C'].width = 32
    else:
        ws.column_dimensions['B'].width = 32

# --- CSV & HTML (.xls) Export Helpers ---
def export_csv(filepath, subjects, records, include_arm=False):
    import csv
    with open(filepath, 'w', newline='', encoding='utf-8') as f:
        writer = csv.writer(f)
        headers = ["Position", "S/N"]
        if include_arm:
            headers.append("Arm")
        headers.append("Student Name")
        headers.extend(subjects)
        headers.extend(["Subjects Offered", "Total Score", "Average Score (%)"])
        writer.writerow(headers)
        
        for r in records:
            row = [r['position'], r['rank']]
            if include_arm:
                row.append(r['arm'])
            row.append(r['name'])
            for s in subjects:
                sc = r['scores'].get(s, 0.0)
                row.append(sc if sc > 0 else "")
            row.extend([r['subjects_offered'], r['total'], r['avg']])
            writer.writerow(row)

def export_html_xls(filepath, title, term_label, subjects, records, include_arm=False):
    """Creates a native HTML table formatted as .xls matching the DLHS portal format."""
    html = f"""<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<title>{title} {term_label}</title>
<style>
  body {{ font-family: Arial, sans-serif; font-size: 11px; }}
  table {{ border-collapse: collapse; width: 100%; }}
  th, td {{ border: 1px solid #999; padding: 4px 6px; text-align: center; }}
  th {{ background-color: #1e3a8a; color: white; }}
  .name {{ text-align: left; font-weight: bold; }}
  .total {{ background-color: #fef3c7; font-weight: bold; }}
  .avg {{ background-color: #e0e7ff; font-weight: bold; }}
  .pos {{ background-color: #dcfce7; font-weight: bold; }}
</style>
</head>
<body>
<h2>DEEPER LIFE HIGH SCHOOL, KADUNA</h2>
<h3>{title} — {term_label} Broadsheet (2025/2026 Session)</h3>
<table border="1">
<thead>
<tr>
  <th>S/N</th>
  {"<th>Arm</th>" if include_arm else ""}
  <th>Student Name</th>
"""
    for sub in subjects:
        html += f"  <th>{sub}</th>\n"
    html += """  <th>Offered</th>
  <th>Total</th>
  <th>Average</th>
  <th>Position</th>
</tr>
</thead>
<tbody>
"""
    for r in records:
        html += f"<tr>\n  <td>{r['rank']}</td>\n"
        if include_arm:
            html += f"  <td>{r['arm']}</td>\n"
        html += f"  <td class='name'>{r['name']}</td>\n"
        for sub in subjects:
            sc = r['scores'].get(sub, 0.0)
            val = sc if sc > 0 else "-"
            html += f"  <td>{val}</td>\n"
        html += f"""  <td>{r['subjects_offered']}</td>
  <td class='total'>{r['total']}</td>
  <td class='avg'>{r['avg']}</td>
  <td class='pos'>{r['position']}</td>
</tr>
"""
    html += """</tbody>
</table>
</body>
</html>
"""
    with open(filepath, 'w', encoding='utf-8') as f:
        f.write(html)

# --- Ensure Directories Exist ---
for d in DEST_DIRS:
    for sub in ['Excel', 'CSV', 'Portal_Exports_XLS']:
        os.makedirs(os.path.join(d, sub), exist_ok=True)

# -----------------------------------------------------------------------------
# 1. GENERATE SS 1 SCIENCE (MERGED A & B)
# -----------------------------------------------------------------------------
print("\n⚙️ Compiling SS 1 Science (Merged Science A & Science B)...")
wb_ss1 = openpyxl.Workbook()
wb_ss1.remove(wb_ss1.active) # Remove default sheet

ss1_term_tables = {}
for term_code, term_label in TERMS:
    subs, recs = prepare_merged_ss1_science(term_code)
    ss1_term_tables[term_code] = (subs, recs)
    
    ws = wb_ss1.create_sheet(title=f"{term_label} (Merged)")
    write_broadsheet_sheet(
        ws,
        title_class="SS 1 SCIENCE (MERGED SCIENCE A & B)",
        term_label=term_label,
        subjects=subs,
        records=recs,
        include_arm=True
    )
    
    # Export individual term CSV & XLS
    for d in DEST_DIRS:
        export_csv(
            os.path.join(d, 'CSV', f"SS1_Science_Merged_{term_code}_2025_2026.csv"),
            subs, recs, include_arm=True
        )
        export_html_xls(
            os.path.join(d, 'Portal_Exports_XLS', f"SS1_Science_Merged_{term_label.replace(' ', '_')}.xls"),
            "SS 1 Science (Merged A & B)", term_label, subs, recs, include_arm=True
        )

# Add Annual Cumulative Sheet to SS 1 Science
ws_annual_ss1 = wb_ss1.create_sheet(title="Annual Cumulative")
write_annual_summary_sheet(
    ws_annual_ss1,
    title_class="SS 1 SCIENCE (MERGED SCIENCE A & B)",
    term_records_dict=ss1_term_tables,
    include_arm=True
)

# Also add standalone SS 1 Science A and SS 1 Science B sheets for convenience
for cid, arm_name, arm_short in [
    (CLASSES['ss1_sci_a']['cid'], "SS 1 SCIENCE A", "Sci A"),
    (CLASSES['ss1_sci_b']['cid'], "SS 1 SCIENCE B", "Sci B")
]:
    for term_code, term_label in TERMS:
        raw_d = fetch_broadsheet(cid, term_code)
        subs_arm, recs_arm = prepare_class_table(raw_d, arm_label=arm_name)
        ws_arm = wb_ss1.create_sheet(title=f"{arm_short} {term_label}")
        write_broadsheet_sheet(
            ws_arm,
            title_class=arm_name,
            term_label=term_label,
            subjects=subs_arm,
            records=recs_arm,
            include_arm=False
        )
        for d in DEST_DIRS:
            export_html_xls(
                os.path.join(d, 'Portal_Exports_XLS', f"{arm_name.replace(' ', '_')}_{term_label.replace(' ', '_')}.xls"),
                arm_name, term_label, subs_arm, recs_arm, include_arm=False
            )

# Save SS1 Workbook
for d in DEST_DIRS:
    wb_ss1.save(os.path.join(d, 'Excel', 'SS1_Science_Merged_2025_2026.xlsx'))
print("  ✓ SS 1 Science Merged Workbook generated.")

# -----------------------------------------------------------------------------
# 2. GENERATE SS 2 HUMANITIES
# -----------------------------------------------------------------------------
print("\n⚙️ Compiling SS 2 Humanities...")
wb_ss2_hum = openpyxl.Workbook()
wb_ss2_hum.remove(wb_ss2_hum.active)

ss2_hum_term_tables = {}
for term_code, term_label in TERMS:
    raw_d = fetch_broadsheet(CLASSES['ss2_hum']['cid'], term_code)
    subs, recs = prepare_class_table(raw_d, arm_label='Humanities')
    ss2_hum_term_tables[term_code] = (subs, recs)
    
    ws = wb_ss2_hum.create_sheet(title=term_label)
    write_broadsheet_sheet(
        ws,
        title_class="SS 2 HUMANITIES",
        term_label=term_label,
        subjects=subs,
        records=recs,
        include_arm=False
    )
    
    for d in DEST_DIRS:
        export_csv(
            os.path.join(d, 'CSV', f"SS2_Humanities_{term_code}_2025_2026.csv"),
            subs, recs, include_arm=False
        )
        export_html_xls(
            os.path.join(d, 'Portal_Exports_XLS', f"SS2_Humanities_{term_label.replace(' ', '_')}.xls"),
            "SS 2 Humanities", term_label, subs, recs, include_arm=False
        )

ws_annual_ss2_hum = wb_ss2_hum.create_sheet(title="Annual Cumulative")
write_annual_summary_sheet(
    ws_annual_ss2_hum,
    title_class="SS 2 HUMANITIES",
    term_records_dict=ss2_hum_term_tables,
    include_arm=False
)

for d in DEST_DIRS:
    wb_ss2_hum.save(os.path.join(d, 'Excel', 'SS2_Humanities_2025_2026.xlsx'))
print("  ✓ SS 2 Humanities Workbook generated.")

# -----------------------------------------------------------------------------
# 3. GENERATE SS 2 SCIENCE
# -----------------------------------------------------------------------------
print("\n⚙️ Compiling SS 2 Science...")
wb_ss2_sci = openpyxl.Workbook()
wb_ss2_sci.remove(wb_ss2_sci.active)

ss2_sci_term_tables = {}
for term_code, term_label in TERMS:
    raw_d = fetch_broadsheet(CLASSES['ss2_sci_a']['cid'], term_code)
    subs, recs = prepare_class_table(raw_d, arm_label='Science A')
    ss2_sci_term_tables[term_code] = (subs, recs)
    
    ws = wb_ss2_sci.create_sheet(title=term_label)
    write_broadsheet_sheet(
        ws,
        title_class="SS 2 SCIENCE (SCIENCE A)",
        term_label=term_label,
        subjects=subs,
        records=recs,
        include_arm=False
    )
    
    for d in DEST_DIRS:
        export_csv(
            os.path.join(d, 'CSV', f"SS2_Science_{term_code}_2025_2026.csv"),
            subs, recs, include_arm=False
        )
        export_html_xls(
            os.path.join(d, 'Portal_Exports_XLS', f"SS2_Science_{term_label.replace(' ', '_')}.xls"),
            "SS 2 Science", term_label, subs, recs, include_arm=False
        )

ws_annual_ss2_sci = wb_ss2_sci.create_sheet(title="Annual Cumulative")
write_annual_summary_sheet(
    ws_annual_ss2_sci,
    title_class="SS 2 SCIENCE (SCIENCE A)",
    term_records_dict=ss2_sci_term_tables,
    include_arm=False
)

for d in DEST_DIRS:
    wb_ss2_sci.save(os.path.join(d, 'Excel', 'SS2_Science_2025_2026.xlsx'))
print("  ✓ SS 2 Science Workbook generated.")

# -----------------------------------------------------------------------------
# 4. GENERATE SS 1 HUMANITIES (Completeness for repeated prompt)
# -----------------------------------------------------------------------------
print("\n⚙️ Compiling SS 1 Humanities (for prompt completeness)...")
wb_ss1_hum = openpyxl.Workbook()
wb_ss1_hum.remove(wb_ss1_hum.active)

ss1_hum_term_tables = {}
for term_code, term_label in TERMS:
    raw_d = fetch_broadsheet(CLASSES['ss1_hum']['cid'], term_code)
    subs, recs = prepare_class_table(raw_d, arm_label='Humanities')
    ss1_hum_term_tables[term_code] = (subs, recs)
    
    ws = wb_ss1_hum.create_sheet(title=term_label)
    write_broadsheet_sheet(
        ws,
        title_class="SS 1 HUMANITIES",
        term_label=term_label,
        subjects=subs,
        records=recs,
        include_arm=False
    )
    
    for d in DEST_DIRS:
        export_csv(
            os.path.join(d, 'CSV', f"SS1_Humanities_{term_code}_2025_2026.csv"),
            subs, recs, include_arm=False
        )
        export_html_xls(
            os.path.join(d, 'Portal_Exports_XLS', f"SS1_Humanities_{term_label.replace(' ', '_')}.xls"),
            "SS 1 Humanities", term_label, subs, recs, include_arm=False
        )

ws_annual_ss1_hum = wb_ss1_hum.create_sheet(title="Annual Cumulative")
write_annual_summary_sheet(
    ws_annual_ss1_hum,
    title_class="SS 1 HUMANITIES",
    term_records_dict=ss1_hum_term_tables,
    include_arm=False
)

for d in DEST_DIRS:
    wb_ss1_hum.save(os.path.join(d, 'Excel', 'SS1_Humanities_2025_2026.xlsx'))
print("  ✓ SS 1 Humanities Workbook generated.")

# -----------------------------------------------------------------------------
# 5. GENERATE MASTER WORKBOOK (ALL CLASSES TOGETHER)
# -----------------------------------------------------------------------------
print("\n⚙️ Compiling Master Broadsheet Workbook (All Classes Combined)...")
wb_master = openpyxl.Workbook()
wb_master.remove(wb_master.active)

# 1. SS1 Sci Merged Term 1, 2, 3, Ann
for t_code, t_lbl in TERMS:
    s, r = ss1_term_tables[t_code]
    ws = wb_master.create_sheet(title=f"SS1 Sci {t_lbl[:4]} (M)")
    write_broadsheet_sheet(ws, "SS 1 SCIENCE (MERGED)", t_lbl, s, r, include_arm=True)

ws = wb_master.create_sheet(title="SS1 Sci Annual")
write_annual_summary_sheet(ws, "SS 1 SCIENCE (MERGED)", ss1_term_tables, include_arm=True)

# 2. SS2 Hum Term 1, 2, 3, Ann
for t_code, t_lbl in TERMS:
    s, r = ss2_hum_term_tables[t_code]
    ws = wb_master.create_sheet(title=f"SS2 Hum {t_lbl[:4]}")
    write_broadsheet_sheet(ws, "SS 2 HUMANITIES", t_lbl, s, r, include_arm=False)

ws = wb_master.create_sheet(title="SS2 Hum Annual")
write_annual_summary_sheet(ws, "SS 2 HUMANITIES", ss2_hum_term_tables, include_arm=False)

# 3. SS2 Sci Term 1, 2, 3, Ann
for t_code, t_lbl in TERMS:
    s, r = ss2_sci_term_tables[t_code]
    ws = wb_master.create_sheet(title=f"SS2 Sci {t_lbl[:4]}")
    write_broadsheet_sheet(ws, "SS 2 SCIENCE", t_lbl, s, r, include_arm=False)

ws = wb_master.create_sheet(title="SS2 Sci Annual")
write_annual_summary_sheet(ws, "SS 2 SCIENCE", ss2_sci_term_tables, include_arm=False)

# 4. SS1 Hum Term 1, 2, 3, Ann
for t_code, t_lbl in TERMS:
    s, r = ss1_hum_term_tables[t_code]
    ws = wb_master.create_sheet(title=f"SS1 Hum {t_lbl[:4]}")
    write_broadsheet_sheet(ws, "SS 1 HUMANITIES", t_lbl, s, r, include_arm=False)

ws = wb_master.create_sheet(title="SS1 Hum Annual")
write_annual_summary_sheet(ws, "SS 1 HUMANITIES", ss1_hum_term_tables, include_arm=False)

for d in DEST_DIRS:
    wb_master.save(os.path.join(d, 'Excel', 'DLHS_Complete_Broadsheets_2025_2026.xlsx'))
print("  ✓ DLHS Master Combined Workbook generated.")

# Save raw JSON backup
raw_json_path = os.path.join(DEST_DIRS[0], 'raw_data_cache.json')
with open(raw_json_path, 'w', encoding='utf-8') as f:
    json.dump(raw_cache, f, indent=2)

print("\n🎉 ALL BROADSHEETS DOWNLOADED, PROCESSED, AND COMPILED SUCCESSFULLY!")
for d in DEST_DIRS:
    print(f"📁 Destination: {d}")
